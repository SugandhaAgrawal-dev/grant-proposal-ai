"""apps/proposals/exporter.py — Export proposal to formatted Word document."""
import io
from datetime import datetime
from django.core.files.base import ContentFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from apps.proposals.models import GrantProposal


COLOR_PRIMARY   = RGBColor(0x00, 0x5A, 0x9C)
COLOR_SECONDARY = RGBColor(0x18, 0x5F, 0xA5)
COLOR_ACCENT    = RGBColor(0x2E, 0x7D, 0x32)


class ProposalExporter:

    def export(self, proposal: GrantProposal) -> None:
        """Build DOCX and save to proposal.docx_file field."""
        doc  = Document()
        self._set_margins(doc)
        self._build_cover(doc, proposal)
        doc.add_page_break()
        self._build_toc(doc, proposal)
        doc.add_page_break()

        for section in proposal.sections.all():
            self._add_section(doc, section.name, section.content)
            doc.add_page_break()

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        filename = f"proposal_{proposal.id}.docx"
        proposal.docx_file.save(filename, ContentFile(buf.read()), save=True)

    def _build_cover(self, doc: Document, proposal: GrantProposal) -> None:
        org_name = proposal.org_profile.name if proposal.org_profile else "Our Organisation"
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(org_name)
        run.font.size  = Pt(28)
        run.font.bold  = True
        run.font.color.rgb = COLOR_PRIMARY

        doc.add_paragraph()
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(proposal.title)
        r.font.size  = Pt(18)
        r.font.color.rgb = COLOR_SECONDARY

        doc.add_paragraph()
        meta = proposal.rfp_metadata or {}
        table = doc.add_table(rows=0, cols=2)
        rows  = [
            ("Submitted by",  org_name),
            ("Agency",        meta.get("agency",   "—")),
            ("Deadline",      meta.get("deadline",  "—")),
            ("Date Prepared", datetime.today().strftime("%B %d, %Y")),
        ]
        for label, value in rows:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = str(value)
            row.cells[0].paragraphs[0].runs[0].font.bold = True

    def _build_toc(self, doc: Document, proposal: GrantProposal) -> None:
        h = doc.add_heading("Table of Contents", level=1)
        if h.runs:
            h.runs[0].font.color.rgb = COLOR_PRIMARY
        for i, sec in enumerate(proposal.sections.all(), 1):
            p = doc.add_paragraph(style="List Number")
            p.add_run(sec.name)

    def _add_section(self, doc: Document, name: str, content: str) -> None:
        h = doc.add_heading(name, level=1)
        if h.runs:
            h.runs[0].font.color.rgb = COLOR_PRIMARY

        for line in content.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("## "):
                sh = doc.add_heading(stripped[3:], level=2)
                if sh.runs:
                    sh.runs[0].font.color.rgb = COLOR_SECONDARY
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith(("- ", "• ", "* ")):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(stripped.lstrip("-•* "))
            else:
                doc.add_paragraph(stripped)

    def _set_margins(self, doc: Document) -> None:
        for section in doc.sections:
            section.top_margin    = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin   = Inches(1.25)
            section.right_margin  = Inches(1.25)
